#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Firmen-Finder für die Strecke Würzburg – Aschaffenburg/Hanau – Frankfurt.

Sucht kleine Betriebe in OpenStreetMap, öffnet die öffentliche Webseite
(inklusive Impressum/Kontakt) und speichert Firma, E-Mail und Geschäftsführer.
Es werden nur öffentliche Seiten gelesen, mit Pause und Beachtung von robots.txt.
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urldefrag, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
USER_AGENT = "FirmenFinder/1.0 (lokale Lead-Recherche; respektiert robots.txt)"
OVERPASS_URLS = (
    "https://overpass-api.de/api/interpreter",
    "https://overpass.kumi.systems/api/interpreter",
    "https://overpass.openstreetmap.fr/api/interpreter",
)

COLUMNS = [
    "Firma",
    "Branche",
    "Geschäftsführer",
    "E-Mail",
    "Weitere E-Mails",
    "Telefon",
    "Straße",
    "PLZ",
    "Ort",
    "Webseite",
    "Quelle",
    "Suchgebiet",
    "Breitengrad",
    "Längengrad",
    "OSM-ID",
    "Status",
]

COLUMN_WIDTHS = {
    "Firma": 36,
    "Branche": 28,
    "Geschäftsführer": 28,
    "E-Mail": 36,
    "Weitere E-Mails": 42,
    "Telefon": 20,
    "Straße": 32,
    "PLZ": 12,
    "Ort": 22,
    "Webseite": 44,
    "Quelle": 44,
    "Suchgebiet": 22,
    "Breitengrad": 14,
    "Längengrad": 14,
    "OSM-ID": 18,
    "Status": 24,
}


@dataclass(frozen=True)
class Ort:
    name: str
    lat: float
    lon: float
    radius: int


# Reihenfolge folgt der Strecke entlang Main / A3.
ORTE = (
    Ort("Würzburg", 49.7945, 9.9290, 7000),
    Ort("Höchberg", 49.7840, 9.8810, 3000),
    Ort("Kist", 49.7430, 9.8460, 3000),
    Ort("Helmstadt", 49.7605, 9.7075, 3000),
    Ort("Marktheidenfeld", 49.8467, 9.6036, 4500),
    Ort("Wertheim", 49.7590, 9.5085, 4500),
    Ort("Miltenberg", 49.7039, 9.2644, 4000),
    Ort("Kleinheubach", 49.7215, 9.2145, 3000),
    Ort("Obernburg am Main", 49.8403, 9.1414, 4000),
    Ort("Elsenfeld", 49.8455, 9.1630, 3000),
    Ort("Aschaffenburg", 49.9770, 9.1521, 6000),
    Ort("Goldbach", 50.0005, 9.1865, 3000),
    Ort("Hösbach", 50.0065, 9.2070, 3500),
    Ort("Kleinostheim", 50.0020, 9.0690, 3000),
    Ort("Mainaschaff", 49.9850, 9.0900, 2500),
    Ort("Stockstadt am Main", 49.9775, 9.0635, 2500),
    Ort("Karlstein am Main", 50.0480, 9.0240, 3500),
    Ort("Alzenau", 50.0883, 9.0647, 4000),
    Ort("Seligenstadt", 50.0436, 8.9753, 4000),
    Ort("Hainburg", 50.0755, 8.9320, 3000),
    Ort("Hanau", 50.1328, 8.9169, 6000),
    Ort("Bruchköbel", 50.1785, 8.9225, 3500),
    Ort("Maintal", 50.1450, 8.8250, 4000),
    Ort("Offenbach am Main", 50.0956, 8.7761, 5500),
    Ort("Frankfurt am Main", 50.1109, 8.6821, 8000),
)

CRAFT_DE = {
    "plumber": "Sanitär / Heizung",
    "electrician": "Elektriker",
    "carpenter": "Zimmerei / Schreinerei",
    "painter": "Maler",
    "roofer": "Dachdecker",
    "hvac": "Heizung / Klima",
    "tiler": "Fliesenleger",
    "gardener": "Garten- und Landschaftsbau",
    "photographer": "Fotograf",
    "tailor": "Schneiderei",
    "shoemaker": "Schuhmacher",
    "watchmaker": "Uhrmacher",
    "jeweller": "Juwelier",
    "bakery": "Bäckerei",
    "butcher": "Metzgerei",
    "locksmith": "Schlüsseldienst",
    "glaziery": "Glaserei",
    "window_construction": "Fensterbau",
    "metal_construction": "Metallbau",
    "blacksmith": "Schmiede",
    "stonemason": "Steinmetz",
    "scaffolder": "Gerüstbauer",
    "plasterer": "Stuckateur",
    "parquet_layer": "Parkettleger",
    "floorer": "Bodenleger",
    "insulation": "Dämmtechnik",
    "solar_panel_installer": "Solarteur",
    "chimney_sweeper": "Schornsteinfeger",
    "chimney_sweep": "Schornsteinfeger",
    "cleaning": "Reinigung",
    "caterer": "Catering",
    "electronics_repair": "Elektronikwerkstatt",
    "computer": "IT-Service",
    "it": "IT-Service",
    "builder": "Bauunternehmen",
    "construction": "Bau",
    "joinery": "Schreinerei",
    "cabinet_maker": "Tischlerei",
    "sawmill": "Sägewerk",
    "upholsterer": "Polsterei",
    "key_cutter": "Schlüsseldienst",
    "optician": "Optiker",
    "dental_technician": "Zahntechnik",
    "beekeeper": "Imkerei",
    "agricultural_engines": "Landmaschinen",
    "car_repair": "Kfz-Werkstatt",
    "motorcycle_repair": "Motorradwerkstatt",
    "bicycle_repair": "Fahrradwerkstatt",
    "boatbuilder": "Bootsbau",
    "winery": "Weingut",
    "brewery": "Brauerei",
    "distillery": "Brennerei",
    "handicraft": "Handwerk",
    "pottery": "Töpferei",
    "sculptor": "Bildhauerei",
    "heating_engineer": "Heizungsbauer",
    "tinsmith": "Klempner",
    "dry_wall": "Trockenbau",
    "signmaker": "Schilder- und Lichtreklame",
    "printer": "Druckerei",
    "bookbinder": "Buchbinderei",
    "window_cleaner": "Fensterreinigung",
    "pest_control": "Schädlingsbekämpfung",
    "sweep": "Schornsteinfeger",
    "photovoltaic": "Photovoltaik",
    "electrician_repair": "Elektro-Werkstatt",
}

SHOP_DE = {
    "car_repair": "Kfz-Werkstatt",
    "car_parts": "Autoteile",
    "tyres": "Reifenservice",
    "motorcycle": "Motorrad",
    "bicycle": "Fahrrad",
    "hairdresser": "Friseur",
    "beauty": "Kosmetik",
    "tattoo": "Tattoo",
    "massage": "Massage",
    "optician": "Optiker",
    "hearing_aids": "Hörgeräte",
    "laundry": "Wäscherei",
    "dry_cleaning": "Reinigung",
    "tailor": "Schneiderei",
    "copyshop": "Copyshop",
    "locksmith": "Schlüsseldienst",
    "paint": "Farben / Malerbedarf",
    "trade": "Fachhandel",
    "glaziery": "Glaserei",
    "hardware": "Eisenwaren",
    "doityourself": "Baumarkt",
    "electronics": "Elektronik",
    "mobile_phone": "Mobilfunk",
    "computer": "Computer",
    "hifi": "HiFi",
    "kitchen": "Küchenstudio",
    "tiles": "Fliesen",
    "flooring": "Bodenbeläge",
    "window": "Fenster",
    "doors": "Türen",
    "lighting": "Leuchten",
    "bathroom_furnishing": "Badstudio",
    "appliance": "Hausgeräte",
    "interior_decoration": "Raumausstattung",
    "furniture": "Möbel",
    "bed": "Betten",
    "sewing": "Nähstudio",
    "watches": "Uhren",
    "jewelry": "Schmuck",
    "photo": "Fotostudio",
    "florist": "Blumen",
    "garden_centre": "Gartenmarkt",
    "bakery": "Bäckerei",
    "butcher": "Metzgerei",
    "pastry": "Konditorei",
    "funeral_directors": "Bestattung",
    "travel_agency": "Reisebüro",
    "estate_agent": "Immobilien",
    "vacuum_cleaner": "Staubsauger-Service",
    "electrical": "Elektro",
    "heating": "Heizung",
    "tool_hire": "Geräteverleih",
    "frame": "Einrahmung",
    "art": "Kunst",
    "craft": "Kunsthandwerk",
    "music": "Musik",
    "musical_instrument": "Musikinstrumente",
    "agrarian": "Agrarbedarf",
}

OFFICE_DE = {
    "accountant": "Buchhaltung",
    "advertising_agency": "Werbung",
    "architect": "Architekturbüro",
    "consulting": "Unternehmensberatung",
    "employment_agency": "Personalvermittlung",
    "engineer": "Ingenieurbüro",
    "estate_agent": "Immobilien",
    "financial_advisor": "Finanzberatung",
    "graphic_design": "Grafikdesign",
    "insurance": "Versicherung",
    "it": "IT-Dienstleister",
    "lawyer": "Rechtsanwaltskanzlei",
    "moving_company": "Umzug",
    "notary": "Notar",
    "property_management": "Hausverwaltung",
    "surveyor": "Vermessung",
    "tax_advisor": "Steuerberatung",
}

AMENITY_DE = {
    "car_repair": "Kfz-Werkstatt",
    "car_wash": "Autowäsche",
    "driving_school": "Fahrschule",
    "vehicle_inspection": "Prüfstelle",
    "workshop": "Werkstatt",
}

EXACT_CHAINS = {
    "dm",
    "o2",
    "kik",
    "nkd",
    "obi",
    "toom",
    "aldi",
    "lidl",
    "rewe",
    "edeka",
    "penny",
    "norma",
    "tegut",
    "ikea",
    "saturn",
    "rossmann",
    "hornbach",
    "bauhaus",
    "fressnapf",
    "fielmann",
}
PREFIX_CHAINS = EXACT_CHAINS | {
    "kaufland",
    "media markt",
    "mediamarkt",
    "apollo optik",
    "deichmann",
    "burger king",
    "mcdonald",
    "mcdonalds",
    "backwerk",
    "subway",
    "starbucks",
    "nordsee",
    "vodafone",
    "dm-drogerie",
    "dm drogerie",
}

SOCIAL_HOSTS = {
    "facebook.com",
    "instagram.com",
    "fb.com",
    "fb.me",
    "twitter.com",
    "x.com",
    "linkedin.com",
    "youtube.com",
    "youtu.be",
    "tiktok.com",
    "google.com",
    "maps.google.com",
    "goo.gl",
    "wikipedia.org",
    "wikidata.org",
    "openstreetmap.org",
    "yelp.com",
    "yelp.de",
    "gelbeseiten.de",
    "dasoertliche.de",
    "11880.com",
    "meinestadt.de",
    "golocal.de",
    "linktr.ee",
    "treatwell.de",
    "jameda.de",
    "doctolib.de",
    "calendly.com",
    "xing.com",
    "pinterest.com",
    "wa.me",
    "t.me",
    "telegram.me",
}

BAD_EMAIL_DOMAINS = {
    "example.com",
    "example.org",
    "example.net",
    "example.de",
    "domain.com",
    "domain.de",
    "domain.tld",
    "email.com",
    "test.de",
    "test.com",
    "invalid",
    "localhost",
    "sentry.io",
    "wixpress.com",
    "schema.org",
    "godaddy.com",
    "wordpress.com",
    "squarespace.com",
    "muster.de",
    "musterfirma.de",
    "ihre-domain.de",
    "ihredomain.de",
    "meinedomain.de",
    "meine-domain.de",
    "ionos.de",
    "ionos.com",
    "strato.de",
    "1und1.de",
    "hubspot.com",
    "hubspotemail.net",
    "hsforms.com",
    "hs-analytics.net",
    "mailchimp.com",
    "list-manage.com",
    "sendgrid.net",
    "salesforce.com",
    "cloudflare.com",
    "cloudflareinsights.com",
    "cookiebot.com",
    "usercentrics.com",
    "usercentrics.eu",
    "onetrust.com",
}

_EMAIL_RE = re.compile(
    r"\b[A-Z0-9][A-Z0-9._%+\-]{0,63}@[A-Z0-9][A-Z0-9.\-]{0,253}\.[A-Z]{2,24}\b",
    re.I,
)
_BAD_TLDS = {"png", "jpg", "jpeg", "gif", "svg", "webp", "css", "js", "ico", "bmp", "woff", "ttf"}
_BAD_LOCAL = {"noreply", "no-reply", "donotreply", "do-not-reply", "mailer-daemon", "bounce"}
_PREFERRED_LOCAL = (
    "info",
    "kontakt",
    "contact",
    "mail",
    "office",
    "service",
    "hallo",
    "hello",
    "anfrage",
    "buero",
    "büro",
    "zentrale",
    "team",
    "firma",
)
_LOW_LOCAL = {
    "datenschutz",
    "privacy",
    "abuse",
    "postmaster",
    "webmaster",
    "support",
    "jobs",
    "karriere",
    "bewerbung",
    "presse",
    "newsletter",
}
_TITLES = {
    "dr",
    "dr.",
    "prof",
    "prof.",
    "dipl.-ing",
    "dipl.-ing.",
    "dipl.-kfm",
    "dipl.-kfm.",
    "ing",
    "ing.",
    "herr",
    "frau",
    "hr",
    "hr.",
    "fr",
    "fr.",
}
_PARTICLES = {"von", "van", "de", "zu", "del", "da", "der", "den", "ten", "di", "dos", "le", "la"}
_NAME_STOP = {
    "gmbh",
    "ag",
    "kg",
    "ohg",
    "ug",
    "gbr",
    "se",
    "co",
    "mbh",
    "haftungsbeschränkt",
    "haftungsbeschraenkt",
    "gesellschaft",
    "firma",
    "unternehmen",
    "straße",
    "strasse",
    "str",
    "weg",
    "platz",
    "allee",
    "gasse",
    "ring",
    "telefon",
    "tel",
    "fax",
    "mobil",
    "mail",
    "email",
    "e-mail",
    "www",
    "http",
    "https",
    "impressum",
    "kontakt",
    "startseite",
    "home",
    "amtsgericht",
    "registergericht",
    "hrb",
    "ust",
    "idnr",
    "steuernummer",
    "deutschland",
    "bayern",
    "hessen",
    "und",
    "oder",
    "sowie",
    "bzw",
    "siehe",
    "angaben",
    "gemäß",
    "gemaess",
    "inhaltlich",
    "verantwortlich",
    "datenschutz",
    "leistungen",
    "unser",
    "unsere",
    "team",
    "über",
    "ueber",
    "uns",
    "mich",
    "montag",
    "dienstag",
    "mittwoch",
    "donnerstag",
    "freitag",
    "samstag",
    "sonntag",
    "uhr",
    "öffnungszeiten",
    "oeffnungszeiten",
    "der",
    "die",
    "das",
    "den",
    "dem",
    "des",
    "ein",
    "eine",
    "einer",
    "welcher",
    "welche",
    "wird",
    "werden",
    "durch",
    "nach",
    "bei",
    "mit",
    "für",
    "fuer",
    "vom",
    "zum",
    "zur",
    "im",
    "in",
    "an",
    "name",
    "inhaber",
    "inhaberin",
    "geschäftsführer",
    "geschäftsführerin",
    "geschaeftsfuehrer",
    "geschaeftsfuehrerin",
}
_TOKEN_RE = re.compile(r"^[A-ZÄÖÜ][A-Za-zÄÖÜäöüß'\-]{1,40}$")
_INITIAL_RE = re.compile(r"^[A-ZÄÖÜ]\.?$")
_RAW_LABELS = (
    "gesetzlich vertreten durch",
    "vertreten durch den geschäftsführer",
    "vertreten durch die geschäftsführerin",
    "vertreten durch den inhaber",
    "vertreten durch die inhaberin",
    "vertreten durch",
    "vertretungsberechtigt",
    "betriebsinhaberin",
    "betriebsinhaber",
    "geschäftsführerin",
    "geschäftsführer",
    "geschaeftsfuehrerin",
    "geschaeftsfuehrer",
    "inhaberin",
    "inhaber",
    "geschäftsleitung",
    "geschaeftsleitung",
    "geschäftsführung",
    "geschaeftsfuehrung",
)
LABELS = tuple(sorted((label.casefold() for label in _RAW_LABELS), key=len, reverse=True))
_LINK_HINTS = (
    "impressum",
    "imprint",
    "anbieterkennzeichnung",
    "kontakt",
    "contact",
    "ueber-uns",
    "über-uns",
    "ueberuns",
    "about",
    "legal",
)
_ORG_TYPES = (
    "Organization",
    "LocalBusiness",
    "Store",
    "WebSite",
    "ProfessionalService",
    "HomeAndConstructionBusiness",
    "AutomotiveBusiness",
    "FoodEstablishment",
)
_PHONE_RE = re.compile(
    r"(?:Telefon|Tel\.?|Fon|Phone)\s*[:.]?\s*(\+?[0-9][0-9()\/\-\s]{6,24})",
    re.I,
)
_GENERIC_NAMES = {
    "werkstatt",
    "praxis",
    "büro",
    "buero",
    "shop",
    "laden",
    "firma",
    "unternehmen",
    "handwerk",
    "service",
    "dienstleistung",
}

_robots_cache: dict[str, RobotFileParser | None] = {}


def configure_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8", errors="replace")


def strip_www(host: str) -> str:
    host = (host or "").lower()
    if host.startswith("www."):
        return host[4:]
    return host


def site_host(url: str) -> str:
    return strip_www(urlparse(url).netloc)


def url_key(url: str) -> str:
    parsed = urlparse(url)
    path = parsed.path.rstrip("/").lower()
    return site_host(url) + path


def same_host(left: str, right: str) -> bool:
    return site_host(left) == site_host(right) and bool(site_host(left))


def is_social(url: str) -> bool:
    host = site_host(url)
    return any(host == domain or host.endswith("." + domain) for domain in SOCIAL_HOSTS)


def normalize_url(raw: str) -> str:
    text = (raw or "").strip()
    if not text:
        return ""
    if text.startswith("//"):
        text = "https:" + text
    if not text.startswith(("http://", "https://")):
        text = "https://" + text
    parsed = urlparse(text)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    return urldefrag(text)[0]


def norm_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def is_chain(name: str, brand: str) -> bool:
    company = norm_space(name).casefold()
    chain_brand = norm_space(brand).casefold()
    if not company and not chain_brand:
        return False
    if company in EXACT_CHAINS or chain_brand in EXACT_CHAINS or chain_brand in PREFIX_CHAINS:
        return True
    return any(company.startswith(chain + " ") or company.startswith(chain + "-") for chain in PREFIX_CHAINS)


def branche_label(tags: dict) -> str:
    groups = (
        ("craft", CRAFT_DE, "Handwerk"),
        ("amenity", AMENITY_DE, "Dienstleister"),
        ("shop", SHOP_DE, "Fachbetrieb"),
        ("office", OFFICE_DE, "Dienstleister"),
    )
    for key, mapping, fallback in groups:
        value = (tags.get(key) or "").strip()
        if not value:
            continue
        if value in {"yes", "no"}:
            return fallback
        return mapping.get(value, f"{fallback} ({value})")
    return "Betrieb"


def pick_website(tags: dict) -> str:
    for key in ("website", "contact:website", "url", "contact:url"):
        raw = tags.get(key) or ""
        for part in raw.split(";"):
            url = normalize_url(part)
            if url and not is_social(url):
                return url
    return ""


def tag_value(tags: dict, *keys: str) -> str:
    for key in keys:
        value = norm_space(tags.get(key) or "")
        if value:
            return value
    return ""


def format_street(tags: dict) -> str:
    street = tag_value(tags, "addr:street")
    number = tag_value(tags, "addr:housenumber")
    if street and number:
        return f"{street} {number}"
    return street


def element_point(element: dict) -> tuple[float | None, float | None]:
    if "lat" in element and "lon" in element:
        return float(element["lat"]), float(element["lon"])
    center = element.get("center") or {}
    if "lat" in center and "lon" in center:
        return float(center["lat"]), float(center["lon"])
    return None, None


def deobfuscate(text: str) -> str:
    cleaned = text or ""
    cleaned = re.sub(r"\s*[\[({]\s*(?:at|ät)\s*[\])}]\s*", "@", cleaned, flags=re.I)
    cleaned = re.sub(r"\s*[\[({]\s*dot\s*[\])}]\s*", ".", cleaned, flags=re.I)
    return cleaned


def valid_email(email: str) -> bool:
    candidate = (email or "").strip().strip(".").lower()
    if not _EMAIL_RE.fullmatch(candidate) or ".." in candidate:
        return False
    local, domain = candidate.split("@", 1)
    base = local.split("+", 1)[0]
    if len(local) < 2 or len(domain) < 4 or base in _BAD_LOCAL or base.startswith("noreply"):
        return False
    tld = domain.rsplit(".", 1)[-1]
    if tld in _BAD_TLDS or tld.isdigit():
        return False
    if domain in BAD_EMAIL_DOMAINS or any(domain.endswith("." + item) for item in BAD_EMAIL_DOMAINS):
        return False
    return True


def find_emails(text: str) -> list[str]:
    found: list[str] = []
    for match in _EMAIL_RE.finditer(deobfuscate(text or "")):
        email = match.group(0).lower().strip(".")
        if valid_email(email):
            found.append(email)
    return found


def choose_emails(emails: Iterable[str], website: str) -> tuple[str, str]:
    host = site_host(website)
    unique: list[str] = []
    seen: set[str] = set()
    for raw in emails:
        email = raw.strip().lower().strip(".")
        if email in seen or not valid_email(email):
            continue
        seen.add(email)
        unique.append(email)

    def score(email: str) -> int:
        local, domain = email.split("@", 1)
        value = 0
        if host and (domain == host or host.endswith("." + domain) or domain.endswith("." + host)):
            value += 100
        base = local.split("+", 1)[0]
        if base in _PREFERRED_LOCAL or any(base.startswith(item) for item in _PREFERRED_LOCAL):
            value += 40
        if base in _LOW_LOCAL:
            value -= 50
        return value

    ranked = sorted(unique, key=score, reverse=True)
    if not ranked:
        return "", ""
    primary = ranked[0]
    further = [email for email in ranked[1:] if score(email) >= 0][:3]
    if score(primary) < 0:
        return primary, ""
    return primary, "; ".join(further)


def normalize_person(raw: str) -> str:
    text = norm_space(raw)
    if not text:
        return ""
    text = re.sub(r"\([^)]*\)", " ", text)
    text = norm_space(text)
    text = re.split(
        r"\||\b(?:Tel(?:efon)?|Fax|Mobil|E-?Mail|Mail|Internet|www\.|https?:|Amtsgericht|Registergericht|HRB|USt-?Id|Straße|Strasse|Str\.)\b",
        text,
        maxsplit=1,
        flags=re.I,
    )[0]
    text = re.split(
        r"\b(?:GmbH(?:\s*&\s*Co\.\s*KG)?|UG(?:\s*\(haftungsbeschränkt\))?|AG|OHG|GbR|e\. ?K\.?|mbH)\b",
        text,
        maxsplit=1,
        flags=re.I,
    )[0]
    text = text.split(",")[0].strip(" .:-–—|/\\")
    text = re.split(r"\s+und\s+", text, maxsplit=1, flags=re.I)[0].strip()
    if not text or len(text) > 70 or re.search(r"\d", text):
        return ""

    tokens = text.split()
    kept_titles: list[str] = []
    while tokens and tokens[0].casefold() in _TITLES:
        title = tokens.pop(0)
        if title.casefold() in {"herr", "frau", "hr", "hr.", "fr", "fr."}:
            continue
        if not title.endswith("."):
            title += "."
        kept_titles.append(title if title[0].isupper() else title.capitalize())

    if not tokens or len(tokens) > 6:
        return ""

    name_tokens = 0
    for token in tokens:
        folded = token.casefold().strip(".")
        if token.casefold() in _PARTICLES or folded in _PARTICLES:
            continue
        if folded in _NAME_STOP:
            return ""
        if _INITIAL_RE.fullmatch(token) or _TOKEN_RE.fullmatch(token):
            name_tokens += 1
            continue
        return ""

    if name_tokens < 2 or name_tokens > 4:
        return ""
    return " ".join(kept_titles + tokens)


def find_manager(text: str) -> str:
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    for index, line in enumerate(lines):
        folded = line.casefold()
        for label in LABELS:
            start = folded.find(label)
            if start < 0:
                continue
            if start > 0 and folded[start - 1].isalnum():
                continue
            after = line[start + len(label) :].strip(" \t:|-–—")
            candidates = [after] if after else []
            if index + 1 < len(lines):
                candidates.append(lines[index + 1])
            for candidate in candidates:
                name = normalize_person(candidate)
                if name:
                    return name
    return ""


def clean_phone(raw: str) -> str:
    phone = norm_space(raw).strip(" .")
    digits = re.sub(r"\D", "", phone)
    if len(digits) < 6:
        return ""
    return phone[:40]


def find_phone(text: str) -> str:
    match = _PHONE_RE.search(text or "")
    if not match:
        return ""
    return clean_phone(match.group(1))


def clean_title(title: str) -> str:
    text = norm_space(title)
    if not text:
        return ""
    text = re.split(r"\s[\|\-–—»«]\s", text, maxsplit=1)[0]
    text = re.sub(r"^(?:Willkommen bei|Herzlich willkommen bei)\s+", "", text, flags=re.I)
    text = re.sub(r"\b(Impressum|Kontakt|Startseite|Home|Willkommen)\b", "", text, flags=re.I)
    return text.strip(" |-–—:")


def choose_company(osm_name: str, page_name: str, website: str) -> str:
    osm_clean = norm_space(osm_name)
    page_clean = norm_space(page_name)
    if osm_clean and osm_clean.casefold() not in _GENERIC_NAMES and len(osm_clean) > 2:
        return osm_clean
    if page_clean and len(page_clean) > 2:
        return page_clean
    if osm_clean:
        return osm_clean
    return site_host(website)


def status_for(email: str, person: str) -> str:
    if email and person:
        return "vollständig"
    if email:
        return "nur E-Mail"
    if person:
        return "nur Geschäftsführer"
    return "ohne Kontaktdaten"


def mailto_emails(soup: BeautifulSoup) -> list[str]:
    found: list[str] = []
    for link in soup.find_all("a", href=True):
        href = str(link["href"]).strip()
        if not href.lower().startswith("mailto:"):
            continue
        raw = unquote(href[7:].split("?", 1)[0]).strip()
        for part in raw.split(","):
            email = part.strip().lower()
            if valid_email(email):
                found.append(email)
    return found


def _walk_jsonld(node: object, acc: dict, depth: int = 0) -> None:
    if depth > 8:
        return
    if isinstance(node, list):
        for item in node:
            _walk_jsonld(item, acc, depth + 1)
        return
    if not isinstance(node, dict):
        return
    kind = node.get("@type", "")
    if isinstance(kind, list):
        kind_text = " ".join(str(item) for item in kind)
    else:
        kind_text = str(kind)
    if any(marker in kind_text for marker in _ORG_TYPES):
        if isinstance(node.get("name"), str):
            acc["names"].append(node["name"])
        if isinstance(node.get("legalName"), str):
            acc["names"].append(node["legalName"])
    for key in ("email", "telephone"):
        value = node.get(key)
        if isinstance(value, str):
            acc[key + "s"].append(value)
    for key in ("founder", "employee", "director"):
        people = node.get(key)
        if people is None:
            continue
        if not isinstance(people, list):
            people = [people]
        for person in people:
            if isinstance(person, str):
                acc["people"].append(person)
            elif isinstance(person, dict) and isinstance(person.get("name"), str):
                acc["people"].append(person["name"])
    for value in node.values():
        if isinstance(value, (dict, list)):
            _walk_jsonld(value, acc, depth + 1)


def jsonld_bits(soup: BeautifulSoup) -> dict:
    acc = {"names": [], "emails": [], "telephones": [], "people": []}
    for script in soup.find_all("script"):
        kind = script.get("type") or ""
        if "ld+json" not in kind.lower():
            continue
        raw = script.string or script.get_text()
        if not raw or not raw.strip():
            continue
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        _walk_jsonld(data, acc)
    return acc


def link_rank(blob: str) -> int:
    if "impressum" in blob or "imprint" in blob or "anbieterkennzeichnung" in blob:
        return 0
    if "kontakt" in blob or "contact" in blob:
        return 1
    return 2


def interesting_links(soup: BeautifulSoup, base_url: str) -> list[str]:
    ranked: list[tuple[int, str]] = []
    for link in soup.find_all("a", href=True):
        href = str(link["href"]).strip()
        if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
            continue
        absolute = urldefrag(urljoin(base_url, href))[0]
        if not absolute.startswith(("http://", "https://")):
            continue
        if not same_host(absolute, base_url) or is_social(absolute):
            continue
        label = link.get_text(" ", strip=True)
        blob = f"{label} {absolute}".casefold()
        if any(hint in blob for hint in _LINK_HINTS):
            ranked.append((link_rank(blob), absolute))
    ranked.sort(key=lambda item: item[0])
    unique: list[str] = []
    seen: set[str] = set()
    for _, url in ranked:
        key = url_key(url)
        if key in seen:
            continue
        seen.add(key)
        unique.append(url)
    parsed = urlparse(base_url)
    root = f"{parsed.scheme}://{parsed.netloc}"
    joined = " ".join(unique).casefold()
    if "impressum" not in joined and "imprint" not in joined:
        for path in ("/impressum", "/impressum.html", "/kontakt"):
            guess = root + path
            key = url_key(guess)
            if key not in seen:
                unique.append(guess)
                seen.add(key)
    return unique[:3]


def first_clean_name(names: Iterable[str]) -> str:
    for name in names:
        cleaned = clean_title(str(name))
        if cleaned and 1 < len(cleaned) <= 120:
            return cleaned
    return ""


def parse_page(html: str, base_url: str) -> dict:
    empty = {"emails": [], "person": "", "page_name": "", "phone": "", "links": []}
    if not html or not html.strip():
        return empty
    try:
        soup = BeautifulSoup(html, "lxml")
    except Exception:
        soup = BeautifulSoup(html, "html.parser")
    emails = mailto_emails(soup)
    linked = jsonld_bits(soup)
    emails.extend(find_emails(" ".join(linked["emails"])))
    links = interesting_links(soup, base_url)
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    og = soup.find("meta", property="og:site_name")
    og_name = og.get("content", "") if og else ""
    page_name = clean_title(str(og_name)) or first_clean_name(linked["names"]) or clean_title(title)
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    emails.extend(find_emails(text))
    person = ""
    for candidate in linked["people"]:
        person = normalize_person(str(candidate))
        if person:
            break
    if not person:
        person = find_manager(text)
    phone = ""
    for raw_phone in linked["telephones"]:
        phone = clean_phone(str(raw_phone))
        if phone:
            break
    if not phone:
        phone = find_phone(text)
    return {
        "emails": emails,
        "person": person,
        "page_name": page_name,
        "phone": phone,
        "links": links,
    }


def decode_html(content: bytes, content_type: str) -> str:
    charset = ""
    lowered = content_type.lower()
    if "charset=" in lowered:
        charset = lowered.split("charset=", 1)[1].split(";", 1)[0].strip(" \"'")
    if charset:
        try:
            return content.decode(charset, errors="replace")
        except LookupError:
            pass
    try:
        from charset_normalizer import from_bytes

        match = from_bytes(content).best()
        if match is not None:
            return str(match)
    except Exception:
        pass
    return content.decode("utf-8", errors="replace")


def origin_of(url: str) -> str:
    parsed = urlparse(url)
    return f"{parsed.scheme}://{parsed.netloc}"


def robots_allowed(session: requests.Session, url: str) -> bool:
    origin = origin_of(url)
    if origin not in _robots_cache:
        parser = RobotFileParser()
        try:
            response = session.get(origin + "/robots.txt", timeout=(6, 12))
            try:
                if response.status_code >= 400:
                    _robots_cache[origin] = None
                else:
                    text = decode_html(response.content[:200_000], response.headers.get("Content-Type", ""))
                    parser.parse(text.splitlines())
                    _robots_cache[origin] = parser
            finally:
                response.close()
        except requests.RequestException:
            _robots_cache[origin] = None
    parser = _robots_cache[origin]
    if parser is None:
        return True
    try:
        return parser.can_fetch(USER_AGENT, url)
    except Exception:
        return True


def fetch_html(session: requests.Session, url: str) -> tuple[str, str] | None:
    if not robots_allowed(session, url):
        logging.info("robots.txt verbietet %s", url)
        return None
    response = None
    try:
        for attempt in range(2):
            try:
                response = session.get(url, timeout=(8, 25), allow_redirects=True, stream=True)
            except requests.RequestException as exc:
                logging.info("Nicht erreichbar: %s (%s)", url, exc.__class__.__name__)
                return None
            if response.status_code in {429, 503} and attempt == 0:
                response.close()
                response = None
                time.sleep(5)
                continue
            break
        if response is None:
            return None
        if response.status_code == 404:
            logging.debug("404 %s", url)
            return None
        if response.status_code >= 400:
            logging.info("HTTP %s %s", response.status_code, url)
            return None
        final_url = response.url
        if is_social(final_url):
            return None
        content_type = response.headers.get("Content-Type", "")
        lowered = content_type.lower()
        if "pdf" in lowered or final_url.lower().split("?", 1)[0].endswith(".pdf"):
            logging.info("PDF übersprungen: %s", final_url)
            return None
        if lowered.startswith(("image/", "video/", "audio/")) or "octet-stream" in lowered:
            return None
        chunks: list[bytes] = []
        total = 0
        for chunk in response.iter_content(chunk_size=65536):
            if not chunk:
                continue
            chunks.append(chunk)
            total += len(chunk)
            if total >= 1_500_000:
                break
        return decode_html(b"".join(chunks), content_type), final_url
    finally:
        if response is not None:
            response.close()


def analyze_site(session: requests.Session, website: str, osm_email: str) -> dict:
    result = {
        "emails": [osm_email] if osm_email else [],
        "person": "",
        "page_name": "",
        "phone": "",
        "source": website,
        "blocked": False,
        "reached": False,
    }
    homepage = fetch_html(session, website)
    links: list[str] = []
    if homepage is None and not robots_allowed(session, website):
        result["blocked"] = True
    if homepage:
        html, final_url = homepage
        result["reached"] = True
        result["source"] = final_url
        parsed = parse_page(html, final_url)
        _merge_page(result, parsed, final_url)
        links = parsed["links"]
    else:
        parsed_root = urlparse(website)
        links = [f"{parsed_root.scheme}://{parsed_root.netloc}/impressum", f"{parsed_root.scheme}://{parsed_root.netloc}/kontakt"]

    seen = {url_key(website)}
    for link in links:
        key = url_key(link)
        if key in seen:
            continue
        seen.add(key)
        if result["person"] and choose_emails(result["emails"], website)[0]:
            break
        time.sleep(0.4)
        page = fetch_html(session, link)
        if not page:
            continue
        html, final_url = page
        result["reached"] = True
        parsed = parse_page(html, final_url)
        _merge_page(result, parsed, final_url)
    return result


def _merge_page(result: dict, parsed: dict, final_url: str) -> None:
    if parsed["emails"]:
        result["emails"].extend(parsed["emails"])
        if result["source"] == "" or "impressum" in final_url.casefold() or "kontakt" in final_url.casefold():
            result["source"] = final_url
    if parsed["person"] and not result["person"]:
        result["person"] = parsed["person"]
        result["source"] = final_url
    if parsed["page_name"] and not result["page_name"]:
        result["page_name"] = parsed["page_name"]
    if parsed["phone"] and not result["phone"]:
        result["phone"] = parsed["phone"]


def overpass_query(ort: Ort, radius: int) -> str:
    shop = "^(" + "|".join(SHOP_DE) + ")$"
    office = "^(" + "|".join(OFFICE_DE) + ")$"
    amenity = "^(" + "|".join(AMENITY_DE) + ")$"
    filters = (
        ('craft', None),
        ("shop", shop),
        ("office", office),
        ("amenity", amenity),
    )
    lines: list[str] = []
    for kind in ("node", "way", "relation"):
        for key, regex in filters:
            around = f"(around:{radius},{ort.lat},{ort.lon})"
            if regex:
                lines.append(f'{kind}["{key}"~"{regex}"]{around};')
            else:
                lines.append(f'{kind}["{key}"]{around};')
    body = "\n  ".join(lines)
    return f"[out:json][timeout:180];\n(\n  {body}\n);\nout center tags;"


def fetch_overpass(query: str) -> dict:
    headers = {"User-Agent": USER_AGENT, "Accept": "application/json"}
    errors: list[str] = []
    for url in OVERPASS_URLS:
        for attempt in range(2):
            try:
                response = requests.post(
                    url,
                    data={"data": query},
                    headers=headers,
                    timeout=(15, 240),
                )
                if response.status_code in {429, 502, 503, 504}:
                    errors.append(f"{url} -> HTTP {response.status_code}")
                    time.sleep(8 * (attempt + 1))
                    continue
                response.raise_for_status()
                payload = response.json()
                if "elements" not in payload:
                    errors.append(f"{url} -> {payload.get('remark', 'keine Elemente')}")
                    break
                return payload
            except (requests.RequestException, ValueError) as exc:
                errors.append(f"{url} -> {exc}")
                time.sleep(4)
    raise RuntimeError("Overpass nicht erreichbar: " + " | ".join(errors[-4:]))


class LeadStore:
    def __init__(self, csv_path: Path, xlsx_path: Path, resume: bool) -> None:
        self.csv_path = csv_path
        self.xlsx_path = xlsx_path
        self.rows: list[dict[str, str]] = []
        self.seen_osm: set[str] = set()
        self.seen_urls: set[str] = set()
        if resume and csv_path.exists() and csv_path.stat().st_size > 0:
            self._load()
        self._file = csv_path.open("a", newline="", encoding="utf-8")
        self._writer = csv.DictWriter(self._file, fieldnames=COLUMNS, delimiter=";")
        if not self.rows and self._file.tell() == 0:
            self._file.write("\ufeff")
            self._writer.writeheader()
            self._file.flush()

    def _load(self) -> None:
        with self.csv_path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, delimiter=";")
            for row in reader:
                if not row:
                    continue
                item = {column: (row.get(column) or "") for column in COLUMNS}
                self.rows.append(item)
                if item["OSM-ID"]:
                    self.seen_osm.add(item["OSM-ID"])
                if item["Webseite"]:
                    self.seen_urls.add(url_key(item["Webseite"]))
        logging.info("Vorhandene Datei geladen: %s Einträge", len(self.rows))

    def known(self, osm_id: str, website: str) -> bool:
        if osm_id and osm_id in self.seen_osm:
            return True
        if website and url_key(website) in self.seen_urls:
            return True
        return False

    def add(self, row: dict[str, str]) -> None:
        item = {column: row.get(column, "") for column in COLUMNS}
        self.rows.append(item)
        if item["OSM-ID"]:
            self.seen_osm.add(item["OSM-ID"])
        if item["Webseite"]:
            self.seen_urls.add(url_key(item["Webseite"]))
        self._writer.writerow(item)
        self._file.flush()
        if len(self.rows) % 10 == 0:
            self.write_xlsx()

    def write_xlsx(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leads"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for column, name in enumerate(COLUMNS, 1):
            cell = sheet.cell(1, column, name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row_index, row in enumerate(self.rows, 2):
            for column, name in enumerate(COLUMNS, 1):
                sheet.cell(row_index, column, row.get(name, ""))
        last_row = max(1, len(self.rows) + 1)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 22
        for name, width in COLUMN_WIDTHS.items():
            index = COLUMNS.index(name) + 1
            sheet.column_dimensions[get_column_letter(index)].width = width
        temporary = self.xlsx_path.with_suffix(".tmp.xlsx")
        try:
            workbook.save(temporary)
            temporary.replace(self.xlsx_path)
        except PermissionError:
            logging.warning("Excel-Datei ist geöffnet und konnte nicht geschrieben werden: %s", self.xlsx_path)
        finally:
            if temporary.exists() and self.xlsx_path.exists():
                try:
                    temporary.unlink()
                except OSError:
                    pass

    def close(self) -> None:
        if not self._file.closed:
            self._file.close()
        self.write_xlsx()


def resolve_orte(names: list[str]) -> list[Ort]:
    if not names:
        return list(ORTE)
    chosen: list[Ort] = []
    seen: set[str] = set()
    known = ", ".join(ort.name for ort in ORTE)
    for raw in names:
        query = raw.strip().casefold()
        if not query:
            continue
        exact = [ort for ort in ORTE if ort.name.casefold() == query]
        hits = exact or [ort for ort in ORTE if ort.name.casefold().startswith(query)]
        if not hits:
            raise SystemExit(f"Unbekannter Ort: {raw}\nBekannte Orte: {known}")
        for ort in hits:
            if ort.name not in seen:
                chosen.append(ort)
                seen.add(ort.name)
    return chosen


def output_paths(stem: str) -> tuple[Path, Path]:
    path = Path(stem)
    if not path.is_absolute():
        path = BASE_DIR / path
    if path.suffix.lower() in {".csv", ".xlsx"}:
        path = path.with_suffix("")
    return path.with_suffix(".csv"), path.with_suffix(".xlsx")


def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[
            logging.FileHandler(BASE_DIR / "firmen_finder.log", encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": USER_AGENT,
            "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "de-DE,de;q=0.9,en;q=0.5",
        }
    )
    session.max_redirects = 6
    return session


def run_scan(args: argparse.Namespace) -> int:
    setup_logging()
    if args.nacht and args.limit == 0:
        args.limit = 40
    orte = resolve_orte(args.orte)
    csv_path, xlsx_path = output_paths(args.ausgabe)
    if args.neu:
        for path in (csv_path, xlsx_path, xlsx_path.with_suffix(".tmp.xlsx")):
            if path.exists():
                path.unlink()
    logging.info("Strecke: %s", " -> ".join(ort.name for ort in orte))
    logging.info("Ausgabe: %s", xlsx_path)
    logging.info("Ausgabe: %s", csv_path)
    if args.nacht:
        logging.info("Nachtlauf: höchstens %s neue Webseiten, vorhandene Einträge bleiben", args.limit)
    store = LeadStore(csv_path, xlsx_path, resume=not args.neu)
    session = build_session()
    stats = {
        "osm": 0,
        "ohne_web": 0,
        "ketten": 0,
        "doppelt": 0,
        "abgerufen": 0,
        "gespeichert": 0,
        "fehler": 0,
    }
    fetched = 0
    stop = False
    try:
        for ort in orte:
            radius = args.radius or ort.radius
            logging.info("Suche %s (Umkreis %s m)", ort.name, radius)
            try:
                payload = fetch_overpass(overpass_query(ort, radius))
            except RuntimeError as exc:
                stats["fehler"] += 1
                logging.error("%s", exc)
                time.sleep(4)
                continue
            elements = payload.get("elements") or []
            logging.info("%s: %s OSM-Treffer", ort.name, len(elements))
            for element in elements:
                stats["osm"] += 1
                tags = element.get("tags") or {}
                if tags.get("abandoned") == "yes":
                    continue
                osm_id = f"{element.get('type', '')}/{element.get('id', '')}"
                website = pick_website(tags)
                name = tag_value(tags, "name", "operator")
                if not website:
                    stats["ohne_web"] += 1
                    continue
                if is_chain(name, tag_value(tags, "brand")):
                    stats["ketten"] += 1
                    continue
                if store.known(osm_id, website):
                    stats["doppelt"] += 1
                    continue
                if args.limit and fetched >= args.limit:
                    logging.info("Limit von %s Webseiten erreicht", args.limit)
                    stop = True
                    break
                if args.nur_suche:
                    fetched += 1
                    logging.info("Gefunden: %s | %s", name or website, website)
                    continue
                fetched += 1
                stats["abgerufen"] += 1
                logging.info("[%s] %s | %s", fetched, name or site_host(website), website)
                try:
                    analyzed = analyze_site(session, website, tag_value(tags, "email", "contact:email"))
                except Exception as exc:
                    stats["fehler"] += 1
                    logging.exception("Fehler bei %s: %s", website, exc)
                    time.sleep(args.delay)
                    continue
                if not analyzed["reached"] and not tag_value(tags, "email", "contact:email"):
                    stats["fehler"] += 1
                    logging.info("Keine abrufbare Seite: %s", website)
                    time.sleep(args.delay)
                    continue
                email, more = choose_emails(analyzed["emails"], website)
                person = analyzed["person"]
                lat, lon = element_point(element)
                city = tag_value(tags, "addr:city", "addr:town", "addr:village") or ort.name
                row = {
                    "Firma": choose_company(name, analyzed["page_name"], website),
                    "Branche": branche_label(tags),
                    "Geschäftsführer": person,
                    "E-Mail": email,
                    "Weitere E-Mails": more,
                    "Telefon": tag_value(tags, "phone", "contact:phone", "mobile", "contact:mobile") or analyzed["phone"],
                    "Straße": format_street(tags),
                    "PLZ": tag_value(tags, "addr:postcode"),
                    "Ort": city,
                    "Webseite": website,
                    "Quelle": analyzed["source"],
                    "Suchgebiet": ort.name,
                    "Breitengrad": f"{lat:.6f}" if lat is not None else "",
                    "Längengrad": f"{lon:.6f}" if lon is not None else "",
                    "OSM-ID": osm_id,
                    "Status": status_for(email, person),
                }
                store.add(row)
                stats["gespeichert"] += 1
                logging.info("  -> %s | %s | %s", row["Firma"], email or "keine E-Mail", person or "kein Geschäftsführer")
                time.sleep(args.delay)
            if stop:
                break
            time.sleep(4)
    except KeyboardInterrupt:
        logging.warning("Abbruch. Bisherige Ergebnisse bleiben gespeichert.")
    finally:
        store.close()
    _log_stats(stats, csv_path, xlsx_path)
    return 0


def _log_stats(stats: dict[str, int], csv_path: Path, xlsx_path: Path) -> None:
    logging.info(
        "Fertig. OSM %s, ohne Webseite %s, Ketten %s, Dubletten %s, abgerufen %s, gespeichert %s, Fehler %s",
        stats["osm"],
        stats["ohne_web"],
        stats["ketten"],
        stats["doppelt"],
        stats["abgerufen"],
        stats["gespeichert"],
        stats["fehler"],
    )
    logging.info("CSV: %s", csv_path)
    logging.info("Excel: %s", xlsx_path)


def run_self_test() -> int:
    checks = [
        (normalize_person("Hans Müller") == "Hans Müller", "Vorname Nachname"),
        (normalize_person("Dr. Anna-Maria Schulz") == "Dr. Anna-Maria Schulz", "Titel und Bindestrich"),
        (normalize_person("Die Gesellschaft") == "", "keine Firma als Person"),
        (normalize_person("Max Mustermann GmbH") == "Max Mustermann", "Rechtsform abtrennen"),
        (normalize_person("Maria von der Leyen") == "Maria von der Leyen", "Namenszusatz"),
        (find_manager("Geschäftsführer: Hans Müller\nTelefon: 0931 12345") == "Hans Müller", "Label in Zeile"),
        (find_manager("Vertreten durch den Geschäftsführer\nPetra Klein") == "Petra Klein", "Name in Folgezeile"),
        ("info@maler-mueller.de" in find_emails("Schreiben Sie an info [at] maler-mueller.de"), "E-Mail [at]"),
        ("info@beispiel.de" in find_emails("info (at) beispiel (dot) de"), "E-Mail (at)/(dot)"),
        (not valid_email("logo@2x.png"), "Dateiendung ist keine E-Mail"),
        (not valid_email("privacy@hubspot.com"), "Tracking-Adresse verwerfen"),
        (is_chain("REWE Markt", ""), "Kette erkennen"),
        (not is_chain("Malerbetrieb Müller", ""), "Handwerker behalten"),
        (choose_emails(["privat@gmail.com", "info@maler.de"], "https://www.maler.de")[0] == "info@maler.de", "eigene Domain bevorzugen"),
    ]
    html = '<html><a href="mailto:buero@firma.de?subject=Hallo">Mail</a><p>Inhaberin: Sabine Wolf</p></html>'
    parsed = parse_page(html, "https://firma.de/")
    checks.append((parsed["emails"] == ["buero@firma.de"], "mailto lesen"))
    checks.append((parsed["person"] == "Sabine Wolf", "Inhaberin im HTML"))
    failed = [label for ok, label in checks if not ok]
    if failed:
        print("Selbsttest fehlgeschlagen: " + ", ".join(failed))
        return 1
    print(f"Selbsttest OK ({len(checks)} Prüfungen)")
    return 0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Sucht kleine Betriebe von Würzburg über Aschaffenburg und Hanau bis Frankfurt "
            "und liest Firma, E-Mail und Geschäftsführer von der öffentlichen Webseite."
        )
    )
    parser.add_argument(
        "--orte",
        default="",
        help="Kommagetrennte Ortsnamen. Ohne Angabe wird die ganze Strecke gescannt.",
    )
    parser.add_argument("--radius", type=int, default=0, help="Umkreis in Metern, ersetzt den Orts-Standard")
    parser.add_argument("--limit", type=int, default=0, help="Höchstens so viele Webseiten abrufen (0 = alle)")
    parser.add_argument("--delay", type=float, default=1.2, help="Pause zwischen Betrieben in Sekunden")
    parser.add_argument("--ausgabe", default="leads", help="Dateiname ohne Endung, im Projektordner")
    parser.add_argument("--neu", action="store_true", help="Vorhandene Ergebnisdatei neu beginnen")
    parser.add_argument(
        "--nacht",
        action="store_true",
        help="Unbeaufsichtigter Lauf: vorhandene Datei fortsetzen und höchstens 40 neue Webseiten lesen",
    )
    parser.add_argument("--nur-suche", action="store_true", help="Nur OpenStreetMap abfragen, Webseiten nicht öffnen")
    parser.add_argument("--self-test", action="store_true", help="Eingebaute Prüfung ohne Netzwerk ausführen")
    args = parser.parse_args(argv)
    args.orte = [part.strip() for part in args.orte.split(",") if part.strip()] if args.orte else []
    if args.radius < 0 or args.limit < 0 or args.delay < 0:
        parser.error("Radius, Limit und Pause dürfen nicht negativ sein.")
    return args


def main(argv: list[str] | None = None) -> int:
    configure_stdio()
    args = parse_args(argv)
    if args.self_test:
        return run_self_test()
    return run_scan(args)


if __name__ == "__main__":
    raise SystemExit(main())
